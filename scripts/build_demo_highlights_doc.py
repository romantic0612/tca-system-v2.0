# -*- coding: utf-8 -*-
"""Build the demo highlights DOCX with stable UTF-8 Chinese text."""

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
REPORT_DIR = ROOT / "reports"
IMG_DIR = REPORT_DIR / "demo_assets_20260526"
OUT = REPORT_DIR / "TCA-System-demo-highlights-20260526.docx"


def apply_font(run, size=10.5, bold=False, color=None):
    run.font.name = "SimSun"
    run._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    run.font.size = Pt(size)
    run.bold = bold
    if color:
        run.font.color.rgb = RGBColor.from_string(color)


def setup_styles(doc):
    normal = doc.styles["Normal"]
    normal.font.name = "SimSun"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.space_after = Pt(6)

    heading_specs = [
        ("Heading 1", 16, "2E74B5"),
        ("Heading 2", 13, "2E74B5"),
        ("Heading 3", 11.5, "1F4D78"),
    ]
    for name, size, color in heading_specs:
        style = doc.styles[name]
        style.font.name = "SimSun"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "宋体")
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = RGBColor.from_string(color)
        style.paragraph_format.space_before = Pt(10)
        style.paragraph_format.space_after = Pt(5)


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    paragraph = cell.paragraphs[0]
    run = paragraph.add_run(str(text))
    apply_font(run, size=9.5, bold=bold, color=color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        shade_cell(table.rows[0].cells[idx], "E8EEF5")
        set_cell_text(table.rows[0].cells[idx], header, bold=True, color="1F4D78")
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], value)
    if widths:
        for row in table.rows:
            for idx, width in enumerate(widths):
                row.cells[idx].width = Inches(width)
    doc.add_paragraph("")


def add_callout(doc, text):
    table = doc.add_table(rows=1, cols=1)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    cell = table.cell(0, 0)
    shade_cell(cell, "F4F6F9")
    set_cell_text(cell, text, color="0B2545")
    doc.add_paragraph("")


def add_image(doc, title, image_name, caption):
    doc.add_heading(title, level=2)
    paragraph = doc.add_paragraph(caption)
    paragraph.paragraph_format.space_after = Pt(6)
    image_path = IMG_DIR / image_name
    pic_paragraph = doc.add_paragraph()
    pic_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    pic_paragraph.add_run().add_picture(str(image_path), width=Inches(6.7))
    cap = doc.add_paragraph("图：" + caption)
    cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
    apply_font(cap.runs[0], size=9, color="6B7280")


def build():
    REPORT_DIR.mkdir(exist_ok=True)
    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.75)
    section.bottom_margin = Inches(0.75)
    section.left_margin = Inches(0.8)
    section.right_margin = Inches(0.8)
    setup_styles(doc)

    title = doc.add_paragraph()
    run = title.add_run("TCA-System 当前演示亮点说明")
    apply_font(run, size=22, bold=True, color="0B2545")

    meta = doc.add_paragraph("版本：v2.0 静态产品演示版｜日期：2026-05-26｜用途：组长/队长演示汇报")
    apply_font(meta.runs[0], size=10, color="6B7280")

    add_callout(
        doc,
        "一句话简介：TCA-System 是面向 K-12 数学学习场景的“教师可控多智能体 AI 辅导系统”。"
        "学生端负责答题和 AI 对话，教师端负责实时观察、诊断和 TCA 组人工干预，"
        "管理端负责学生分组、批量导入和实验数据导出。",
    )

    paragraph = doc.add_paragraph(
        "当前版本已经可以用于演示“学生学习过程 -> 系统识别困难 -> 教师端实时看到 -> "
        "TCA 教师 Override -> 学生端接收指导 -> 管理端导出实验数据”的完整闭环。"
    )
    apply_font(paragraph.runs[0], bold=True)

    doc.add_heading("1. 当前可展示的核心亮点", level=1)
    add_table(
        doc,
        ["亮点", "现在能展示什么", "演示时怎么看"],
        [
            ["四组实验条件", "SA / EXP / AI-AUTO / TCA 四组学生均从数据库读取并在教师端区分显示。", "教师端学生列表显示组别、当前 Agent、题号、最后消息、上次活动。"],
            ["智能评估与触发", "学生答题/聊天后，系统可记录错误类型、错误累计，并根据 L1/L2/L3 触发提醒。", "学生说“我不会”等求助词后，教师端出现待处理请求和诊断卡片。"],
            ["TCA 教师 Override", "TCA 组允许教师发送提示、切换 Tutor；学生端可实时收到教师指导。", "教师端选中王芳/TCA，干预按钮可用；学生端顶部显示教师指导入口。"],
            ["权限边界", "非 TCA 组只能观察，不允许教师干预；前端禁用，后端也拒绝。", "教师端选中张明/SA，会显示“仅观察，不可干预”。"],
            ["管理端数据能力", "批量导入有行级错误提示，支持导出分组和实验核心数据。", "管理端有“批量导入”“导出实验数据 (CSV)”入口。"],
            ["部署能力", "Docker + GitHub 主线可部署到云服务器，默认运行可演示的 static 产品页。", "服务器执行 git pull + docker compose up -d --build 即可更新。"],
        ],
        widths=[1.3, 3.0, 2.6],
    )

    doc.add_heading("2. 推荐演示路径", level=1)
    steps = [
        "登录学生端，使用 TCA 学生账号 20240003 / 123456，展示答题、AI 对话和教师指导入口。",
        "登录教师端，使用教师账号 100001 / teacher123，展示待处理请求、学生状态列表、诊断卡片和 TCA 干预模板。",
        "教师端切到非 TCA 学生，例如 20240001 张明，展示“仅观察，不可干预”的权限边界。",
        "登录管理端，使用管理员账号 900001 / admin123，展示学生分组、批量导入、导出实验数据。",
    ]
    for index, step in enumerate(steps, 1):
        doc.add_paragraph(f"{index}. {step}")

    add_image(doc, "3. 学生端：答题、AI 对话与教师指导入口", "01_student_tca_override.png", "学生端展示题目区、AI 对话区、TCA 组教师指导入口，以及底部消息输入。")
    add_image(doc, "4. 教师端：TCA 组诊断与可干预面板", "02_teacher_tca_dashboard.png", "教师端选中 TCA 学生后，可看到待处理请求、学生当前题详情、诊断卡片、提示模板和发送按钮。")
    add_image(doc, "5. 教师端：非 TCA 组仅观察边界", "03_teacher_observe_only.png", "教师端选中非 TCA 学生后，干预区禁用，并明确提示“仅观察，不可干预”。")
    add_image(doc, "6. 管理端：分组、导入与数据导出", "04_admin_export_import.png", "管理端从数据库读取学生分组，可批量导入学生，并导出分组结果和实验数据 CSV。")

    doc.add_heading("7. 目前可验收功能清单", level=1)
    add_table(
        doc,
        ["验收点", "状态", "说明"],
        [
            ["学生端对话和题目切换", "可演示", "对话按题号绑定，聊天区在框内滚动。"],
            ["教师端待处理提醒", "可演示", "学生求助或触发规则后，教师端显示待处理请求和通知条。"],
            ["教师端详情页", "可演示", "按题展示对话、诊断、评分、错误类型、错误累计和触发层级。"],
            ["TCA Override", "可演示", "TCA 组可发送提示/切换 Tutor，学生端实时收到并持久化。"],
            ["非 TCA 权限限制", "已补齐", "前端禁用，后端接口也返回拒绝。"],
            ["管理端批量导入", "已补齐", "返回成功数、失败数和失败行原因。"],
            ["实验数据导出", "已补齐", "导出学生、对话、评估、触发、教师干预记录。"],
            ["题库数据库化", "暂缓", "当前仍使用默认内置题库，符合本阶段计划。"],
        ],
        widths=[1.8, 1.0, 4.2],
    )

    doc.add_heading("8. 本轮关键修改文件", level=1)
    add_table(
        doc,
        ["文件", "改动摘要"],
        [
            ["backend/api/teacher.py", "新增非 TCA 学生后端干预拒绝，保护教师 Override 权限边界。"],
            ["backend/api/student.py", "学生列表接口增加当前 Agent、待处理数、最后消息、上次活动时间。"],
            ["backend/api/admin.py", "增强 CSV 导入校验，新增 experiment-data.csv 实验数据导出接口。"],
            ["frontend/static/teacher.html", "新增通知条、学生列表状态增强、TCA/非TCA干预权限显示、模板干预 UI。"],
            ["frontend/static/admin.html", "导入结果展示失败行原因，新增实验数据导出按钮。"],
            ["test_product_flows.py", "新增产品链路测试：学生求助、教师干预、权限和导入导出。"],
        ],
        widths=[2.2, 4.8],
    )

    add_callout(
        doc,
        "建议汇报口径：当前系统已经具备“能跑、能演示、能验收”的主链路；"
        "下一阶段重点是题库数据库化、前端 Vue3 组件化重构、模板库精细化和更正式的多人并发数据存储。",
    )

    for section in doc.sections:
        footer = section.footer.paragraphs[0]
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.text = "TCA-System · 当前演示亮点说明 · 2026-05-26"
        apply_font(footer.runs[0], size=8.5, color="6B7280")

    doc.save(OUT)
    print(OUT)


if __name__ == "__main__":
    build()
